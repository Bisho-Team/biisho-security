from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE, MSO_CONNECTOR
from pptx.enum.dml import MSO_THEME_COLOR
from pptx.enum.text import MSO_AUTO_SIZE
import os, requests

OUT='/vercel/share/v0-project/Biisho_Company_Pitch_Deck.pptx'
LOGO='/vercel/share/v0-project/Biisho Companylogo - Copy.png'
URL='https://hebbkx1anhila5yf.public.blob.vercel-storage.com/logo%20-%20Copy-cJJzC5GbtfTCGwiMegh7B4jUW678Cu.png'
if not os.path.exists(LOGO):
    r=requests.get(URL, timeout=30); r.raise_for_status(); open(LOGO,'wb').write(r.content)

prs=Presentation(); prs.slide_width=Inches(13.333); prs.slide_height=Inches(7.5)
W,H=13.333,7.5
NAVY=RGBColor(15,29,48); INK=RGBColor(28,42,58); ORANGE=RGBColor(242,104,31); PEACH=RGBColor(255,238,228); SLATE=RGBColor(91,108,125); PALE=RGBColor(247,249,251); WHITE=RGBColor(255,255,255); GREEN=RGBColor(33,145,103); RED=RGBColor(204,68,62); AMBER=RGBColor(226,150,47); LINE=RGBColor(222,228,234)
FONT='Aptos'; DISPLAY='Aptos Display'

def rect(slide,x,y,w,h,fill,line=None,radius=False):
    shp=slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE if radius else MSO_SHAPE.RECTANGLE, Inches(x),Inches(y),Inches(w),Inches(h))
    shp.fill.solid(); shp.fill.fore_color.rgb=fill
    shp.line.color.rgb=line or fill
    if radius: shp.adjustments[0]=0.12
    return shp

def txt(slide,text,x,y,w,h,size=18,color=INK,bold=False,font=FONT,align=PP_ALIGN.LEFT,margin=0.06):
    box=slide.shapes.add_textbox(Inches(x),Inches(y),Inches(w),Inches(h)); tf=box.text_frame; tf.clear(); tf.word_wrap=True
    tf.margin_left=tf.margin_right=Inches(margin); tf.margin_top=tf.margin_bottom=Inches(margin)
    p=tf.paragraphs[0]; p.alignment=align; p.space_after=Pt(0)
    r=p.add_run(); r.text=text; r.font.name=font; r.font.size=Pt(size); r.font.bold=bold; r.font.color.rgb=color
    return box

def add_logo(slide,x=0.55,y=0.35,w=1.35):
    slide.shapes.add_picture(LOGO, Inches(x), Inches(y), width=Inches(w))

def footer(slide,num,dark=False):
    c=RGBColor(205,215,224) if dark else SLATE
    txt(slide,'BIISHO  /  SMART SECURITY OPERATIONS',0.55,7.15,5,0.18,7,c,True)
    txt(slide,f'{num:02d}',12.25,7.12,0.55,0.22,9,c,True,align=PP_ALIGN.RIGHT)

def title(slide,kicker,headline,sub=None,dark=False):
    c=WHITE if dark else NAVY; sc=RGBColor(220,230,238) if dark else SLATE
    txt(slide,kicker.upper(),0.65,0.55,5,0.25,9,ORANGE,True)
    txt(slide,headline,0.65,0.9,11.5,0.7,29,c,True,DISPLAY)
    if sub: txt(slide,sub,0.67,1.68,10.8,0.38,12,sc)

def bullet(slide,text,x,y,w,color=INK,size=14):
    txt(slide,'•',x,y,0.18,0.25,size+2,ORANGE,True)
    txt(slide,text,x+0.25,y,w-0.25,0.42,size,color)

def pill(slide,label,x,y,w,fill=PEACH,color=ORANGE):
    rect(slide,x,y,w,0.34,fill,fill,True); txt(slide,label,x,y+0.03,w,0.22,9,color,True,align=PP_ALIGN.CENTER)

def card(slide,x,y,w,h,heading,body,accent=ORANGE):
    rect(slide,x,y,w,h,WHITE,LINE,True); rect(slide,x,y,0.07,h,accent,accent)
    txt(slide,heading,x+0.25,y+0.22,w-0.4,0.28,15,NAVY,True)
    txt(slide,body,x+0.25,y+0.62,w-0.45,h-0.72,11,SLATE)

def add_node(slide,x,y,w,h,label,fill=WHITE,stroke=LINE,txtcolor=INK):
    rect(slide,x,y,w,h,fill,stroke,True); txt(slide,label,x+0.08,y+0.16,w-0.16,h-0.2,11,txtcolor,True,align=PP_ALIGN.CENTER)

def arrow(slide,x1,y1,x2,y2,color=ORANGE,width=2):
    line=slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(x1),Inches(y1),Inches(x2),Inches(y2)); line.line.color.rgb=color; line.line.width=Pt(width); line.line.end_arrowhead=True

# 1 cover
s=prs.slides.add_slide(prs.slide_layouts[6]); rect(s,0,0,W,H,NAVY); rect(s,0,6.82,W,0.68,ORANGE)
add_logo(s,0.75,0.55,2.0); pill(s,'CLIENT & STAKEHOLDER DECK',0.8,2.05,2.2,RGBColor(52,66,83),RGBColor(255,188,151))
txt(s,'From alert\nto action.',0.75,2.55,7.2,1.5,39,WHITE,True,DISPLAY); txt(s,'Smart security operations for a safer, more accountable world.',0.8,4.35,6.2,0.5,17,RGBColor(216,226,235))
# abstract signal graphic
for i,(x,y,r,c) in enumerate([(9.55,2.1,2.4,RGBColor(43,62,82)),(9.55,2.1,1.75,RGBColor(55,76,98)),(9.55,2.1,1.1,ORANGE)]):
    sh=s.shapes.add_shape(MSO_SHAPE.OVAL,Inches(x-r/2),Inches(y-r/2),Inches(r),Inches(r)); sh.fill.background(); sh.line.color.rgb=c; sh.line.width=Pt(2)
rect(s,9.27,1.82,0.56,0.56,ORANGE,ORANGE,True); txt(s,'B',9.27,1.87,0.56,0.4,23,WHITE,True,align=PP_ALIGN.CENTER)
txt(s,'BIISHO COMPANY',0.8,7.02,3,0.2,10,WHITE,True)

# 2 challenge
s=prs.slides.add_slide(prs.slide_layouts[6]); rect(s,0,0,W,H,PALE); add_logo(s); title(s,'01  /  THE OPPORTUNITY','Security operations should not depend on scattered updates.','When an incident happens, every minute and every handoff matters.');
card(s,0.7,2.35,3.75,2.25,'Fragmented communication','Calls, messages, spreadsheets, and manual follow-up make it harder to see the complete picture.',RED)
card(s,4.8,2.35,3.75,2.25,'Limited visibility','Teams need a shared view of locations, people, severity, ownership, and current status.',AMBER)
card(s,8.9,2.35,3.75,2.25,'Unclear accountability','Without a structured record, it is difficult to prove what happened, who acted, and what improved.',ORANGE)
txt(s,'The result: avoidable delay, duplicated effort, and lower confidence for clients and leaders.',0.75,5.45,11.5,0.4,18,NAVY,True,DISPLAY); footer(s,2)

# 3 answer
s=prs.slides.add_slide(prs.slide_layouts[6]); rect(s,0,0,W,H,WHITE); add_logo(s); title(s,'02  /  THE BIISHO ANSWER','One connected operating picture.','Biisho brings the people, places, incidents, and decisions involved in a response process into one controlled environment.');
# central diagram
add_node(s,0.85,3.0,2.0,0.72,'Protected\nlocations',PEACH,ORANGE,ORANGE); add_node(s,3.3,3.0,2.0,0.72,'Live alarm\nrecord',NAVY,NAVY,WHITE); add_node(s,5.75,3.0,2.0,0.72,'Response\nteam',PEACH,ORANGE,ORANGE); add_node(s,8.2,3.0,2.0,0.72,'Resolution\nrecord',NAVY,NAVY,WHITE); add_node(s,10.65,3.0,1.8,0.72,'Insight &\nlearning',PEACH,ORANGE,ORANGE)
for x in [2.85,5.3,7.75,10.2]: arrow(s,x,3.36,x+0.42,3.36)
pill(s,'DETECT',1.08,4.15,1.5); pill(s,'COORDINATE',3.48,4.15,1.65); pill(s,'RESPOND',5.98,4.15,1.55); pill(s,'CLOSE THE LOOP',8.38,4.15,1.85); pill(s,'IMPROVE',10.78,4.15,1.45)
txt(s,'A practical workflow that turns security events into coordinated action and reliable records.',1.0,5.4,11.2,0.42,18,NAVY,True,DISPLAY,align=PP_ALIGN.CENTER); footer(s,3)

# 4 how it works
s=prs.slides.add_slide(prs.slide_layouts[6]); rect(s,0,0,W,H,PALE); add_logo(s); title(s,'03  /  HOW IT WORKS','From first alert to final outcome.','Every incident follows a clear path, with the right context available at each step.');
steps=[('01','Create alarm','Capture type, severity, location, branch, and notes.'),('02','Notify team','The incident becomes visible to permitted response teams.'),('03','Acknowledge','A responder takes ownership and records progress.'),('04','Resolve','The outcome is documented for reporting and learning.')]
for i,(n,h,b) in enumerate(steps):
 x=0.8+i*3.05; rect(s,x,2.55,2.55,2.05,WHITE,LINE,True); rect(s,x+0.22,2.78,0.52,0.52,ORANGE,ORANGE,True); txt(s,n,x+0.22,2.9,0.52,0.2,11,WHITE,True,align=PP_ALIGN.CENTER); txt(s,h,x+0.22,3.55,2.0,0.28,15,NAVY,True); txt(s,b,x+0.22,3.95,2.05,0.5,11,SLATE)
 if i<3: arrow(s,x+2.58,3.56,x+2.95,3.56,ORANGE,2)
txt(s,'Clear statuses. Clear ownership. A complete record.',0.8,5.45,11.6,0.4,20,ORANGE,True,DISPLAY,align=PP_ALIGN.CENTER); footer(s,4)

# 5 roles
s=prs.slides.add_slide(prs.slide_layouts[6]); rect(s,0,0,W,H,WHITE); add_logo(s); title(s,'04  /  ROLE-BASED EXPERIENCE','Focused workspaces for every responsibility.','Biisho keeps each user focused on the decisions and actions relevant to their role.');
roles=[('Super Administrator','Platform-wide governance, companies, users, settings, alarms, and analytics.'),('Office Manager','Branches, office operators, alarms, and internal operational reporting.'),('Police Manager','Response queue, police operators, pending work, and performance oversight.'),('Office Operator','Front-line monitoring and immediate escalation for permitted locations.'),('Police Operator','Assigned response reports, progress updates, and resolution details.')]
for i,(h,b) in enumerate(roles):
 x=0.75+(i%3)*4.15; y=2.45+(i//3)*1.8; card(s,x,y,3.7,1.35,h,b,[ORANGE,NAVY,ORANGE,AMBER,NAVY][i])
footer(s,5)

# 6 capabilities
s=prs.slides.add_slide(prs.slide_layouts[6]); rect(s,0,0,W,H,PALE); add_logo(s); title(s,'05  /  CORE CAPABILITIES','Built for the rhythm of real operations.','A connected set of capabilities supports daily readiness, active response, and management review.');
cap=[('Live dashboards','Current alarms, workload, response status, activity, and connection health.'),('Alarm management','Create, search, filter, review, and track incidents by status and severity.'),('Location management','Organize companies, branches, contacts, addresses, and optional coordinates.'),('Response tracking','Assign reports, capture acknowledgement, progress, completion, and notes.'),('Analytics & reports','See trends, volumes, workload, performance, and exportable summaries.'),('Account controls','Manage profiles, passwords, access, and secure sign-out.')]
for i,(h,b) in enumerate(cap):
 x=0.75+(i%3)*4.15; y=2.35+(i//3)*1.7; card(s,x,y,3.7,1.25,h,b,[ORANGE,NAVY,AMBER,ORANGE,NAVY,AMBER][i])
footer(s,6)

# 7 dashboard visual
s=prs.slides.add_slide(prs.slide_layouts[6]); rect(s,0,0,W,H,NAVY); add_logo(s); title(s,'06  /  OPERATIONAL VISIBILITY','See what needs attention now.','Role-specific dashboards turn operational data into a shared view of risk, workload, and performance.',True)
# mini dashboard
rect(s,0.75,2.35,11.85,3.55,RGBColor(27,44,64),RGBColor(58,78,99),True)
for i,(lab,val,col) in enumerate([('Active alarms','12',ORANGE),('Pending reports','07',AMBER),('Resolved today','24',GREEN),('Avg. response','08m',RGBColor(128,192,255))]):
 x=1.1+i*2.82; rect(s,x,2.75,2.35,0.9,RGBColor(36,54,75),RGBColor(62,82,103),True); txt(s,lab,x+0.16,2.92,2.0,0.18,9,RGBColor(190,204,216),True); txt(s,val,x+0.16,3.18,2.0,0.3,23,col,True,DISPLAY)
# bars
for i,h in enumerate([0.65,1.15,0.85,1.7,1.35,2.1,1.55]): rect(s,1.25+i*0.48,5.2-h,0.26,h,ORANGE if i==5 else RGBColor(86,116,143),RGBColor(86,116,143),True)
txt(s,'Incident volume / last 7 days',1.25,5.38,3,0.2,9,RGBColor(190,204,216),True)
# status list
for i,(lab,st,col) in enumerate([('Main office / intrusion alert','Active',RED),('North branch / access issue','Acknowledged',AMBER),('Warehouse / false alarm','Resolved',GREEN)]):
 y=4.05+i*0.45; txt(s,lab,5.2,y,4.4,0.2,10,WHITE); pill(s,st,10.2,y-0.04,1.25,RGBColor(47,64,83),col)
footer(s,7,True)

# 8 benefits
s=prs.slides.add_slide(prs.slide_layouts[6]); rect(s,0,0,W,H,WHITE); add_logo(s); title(s,'07  /  BUSINESS VALUE','Confidence that compounds over time.','Biisho makes the response process faster, clearer, and easier to improve.');
benefits=[('Faster response','Real-time alerts and clear ownership reduce the distance between detection and action.'),('Stronger accountability','Time-stamped records show what happened, who acted, and how it ended.'),('Better oversight','Location, severity, status, and personnel context are visible together.'),('Scalable operations','A company-and-branch structure supports multiple organizations and sites.'),('Less administration','Structured records and reports reduce manual logs and repetitive follow-up.'),('Continuous improvement','Analytics reveal recurring issues, workload patterns, and response gaps.')]
for i,(h,b) in enumerate(benefits):
 x=0.75+(i%2)*6.15; y=2.35+(i//2)*1.38; txt(s,f'0{i+1}',x,y,0.48,0.3,14,ORANGE,True); txt(s,h,x+0.72,y,4.9,0.25,15,NAVY,True); txt(s,b,x+0.72,y+0.36,5.1,0.42,11,SLATE)
footer(s,8)

# 9 stakeholder value
s=prs.slides.add_slide(prs.slide_layouts[6]); rect(s,0,0,W,H,PALE); add_logo(s); title(s,'08  /  STAKEHOLDER VALUE','One platform. Multiple wins.','The same operating model creates value for clients, leaders, managers, and front-line responders.');
rows=[('Clients & site owners','Confidence that incidents are visible, assigned, tracked, and documented.'),('Security leadership','A centralized view of risk, locations, people, activity, and trends.'),('Office managers','Practical control of branches, operators, alarms, and internal reports.'),('Police managers','A prioritized response queue with ownership and performance context.'),('Operators','Clear tasks, relevant location details, and a simple path to update status.'),('Senior management','Evidence for service quality, resource planning, and improvement.')]
for i,(h,b) in enumerate(rows):
 y=2.25+i*0.66; rect(s,0.8,y,11.75,0.52,WHITE,LINE,True); txt(s,h,1.05,y+0.14,2.5,0.2,11,NAVY,True); txt(s,b,3.75,y+0.14,8.25,0.2,11,SLATE)
footer(s,9)

# 10 differentiators
s=prs.slides.add_slide(prs.slide_layouts[6]); rect(s,0,0,W,H,NAVY); add_logo(s); title(s,'09  /  WHY BIISHO','Designed around the moments that matter.','Biisho is more than an information system: it is an operating model for responsible response.',True)
for i,(h,b) in enumerate([('Connected by design','The organization, location, alarm, response team, and outcome stay linked.'),('Actionable by default','The workflow is built to move incidents forward, not just display data.'),('Clear at every level','Front-line users get detail; leaders get summaries, trends, and evidence.'),('Built for trust','Access controls, validation, and structured records support responsible operations.')]):
 y=2.35+i*0.88; rect(s,0.85,y,0.12,0.62,ORANGE,ORANGE,True); txt(s,h,1.25,y,3.15,0.25,15,WHITE,True); txt(s,b,4.65,y,7.1,0.33,12,RGBColor(203,216,227))
footer(s,10,True)

# 11 rollout
s=prs.slides.add_slide(prs.slide_layouts[6]); rect(s,0,0,W,H,WHITE); add_logo(s); title(s,'10  /  PATH TO VALUE','Start focused. Scale with confidence.','A practical rollout aligns the platform with the client’s locations, roles, procedures, and success measures.');
for i,(h,b) in enumerate([('Scope','Confirm priority companies, branches, users, and response stakeholders.'),('Configure','Set severity definitions, ownership rules, permissions, and reporting needs.'),('Activate','Onboard teams, establish operating routines, and begin live incident tracking.'),('Improve','Review performance, learn from trends, and expand coverage over time.')]):
 x=0.8+i*3.05; rect(s,x,2.7,2.55,2.1,PALE,LINE,True); rect(s,x,2.7,2.55,0.1,ORANGE,ORANGE); txt(s,f'0{i+1}',x+0.22,3.02,0.45,0.25,15,ORANGE,True); txt(s,h,x+0.22,3.48,2.0,0.25,16,NAVY,True); txt(s,b,x+0.22,3.9,2.05,0.55,11,SLATE)
 if i<3: arrow(s,x+2.58,3.75,x+2.98,3.75,ORANGE,2)
txt(s,'The goal is not just better software. It is better readiness, response, and trust.',0.85,5.55,11.6,0.4,19,NAVY,True,DISPLAY,align=PP_ALIGN.CENTER); footer(s,11)

# 12 close
s=prs.slides.add_slide(prs.slide_layouts[6]); rect(s,0,0,W,H,NAVY); rect(s,0,6.65,W,0.85,ORANGE); add_logo(s,0.75,0.55,2.0); txt(s,'The promise is simple.',0.8,2.0,7.0,0.5,18,RGBColor(255,188,151),True); txt(s,'Help the right people\nsee the right incident,\nact with the right context,\nand close the loop.',0.8,2.6,8.6,2.2,30,WHITE,True,DISPLAY); txt(s,'Biisho Company  /  Smart security operations',0.8,5.55,5.8,0.3,13,RGBColor(208,220,230)); txt(s,'BIISHO',10.35,5.55,2.2,0.3,18,ORANGE,True,align=PP_ALIGN.RIGHT)
footer(s,12,True)

prs.core_properties.title='Biisho Company Pitch Deck'; prs.core_properties.subject='Smart Security Operations Platform'; prs.core_properties.author='Biisho Company'
prs.save(OUT)
print(OUT)
